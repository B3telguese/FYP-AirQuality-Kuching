"""
Air Quality PM2.5 Prediction Dashboard
FYP — Muhamad Hafizzuddin — CDCS2303A
Stacking Ensemble: RF + XGBoost + CatBoost + SVR (base) + XGBoost (meta)
"""
from flask import Flask, render_template, request, jsonify, send_file
import os, io, warnings
from datetime import datetime
import pandas as pd
import numpy as np
import joblib
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from collections import Counter
from sklearn.metrics import mean_squared_error, mean_absolute_error, r2_score

warnings.filterwarnings('ignore')

app = Flask(__name__)
BASE_DIR  = os.path.dirname(os.path.abspath(__file__))
MODEL_DIR = os.path.join(BASE_DIR, 'saved_model')

MODEL = SCALER = ALL_FEATURES = None
COMPARISON = None

REQUIRED_COLS = ['NO2', 'O3', 'CO', 'PM10', 'PM2.5']

# EPA PM2.5 → AQI breakpoints (mirrors notebook Phase 9)
AQI_BP = [
    (0.0,   12.0,   0,   50),
    (12.1,  35.4,  51,  100),
    (35.5,  55.4, 101,  150),
    (55.5, 150.4, 151,  200),
    (150.5, 250.4, 201, 300),
    (250.5, 500.4, 301, 500),
]

CAT_COLORS = {
    'Good':                           '#2ea88a',
    'Moderate':                       '#e3a02a',
    'Unhealthy for Sensitive Groups': '#d4763b',
    'Unhealthy':                      '#e05252',
    'Very Unhealthy':                 '#b75de0',
    'Hazardous':                      '#9e2a2a',
}


def _load_model():
    global MODEL, SCALER, ALL_FEATURES
    files = ('stacking_model.pkl', 'scaler.pkl', 'all_features.pkl')
    if not all(os.path.isfile(os.path.join(MODEL_DIR, f)) for f in files):
        return
    try:
        MODEL        = joblib.load(os.path.join(MODEL_DIR, 'stacking_model.pkl'))
        SCALER       = joblib.load(os.path.join(MODEL_DIR, 'scaler.pkl'))
        ALL_FEATURES = joblib.load(os.path.join(MODEL_DIR, 'all_features.pkl'))
    except Exception:
        pass


_load_model()


def _load_comparison():
    global COMPARISON
    xlsx = os.path.join(BASE_DIR, 'Kuching_model_results.xlsx')
    if not os.path.isfile(xlsx):
        return
    try:
        def _clean(df):
            return df.replace({float('nan'): None}).to_dict('records')

        sheets = pd.ExcelFile(xlsx).sheet_names
        cmp  = _clean(pd.read_excel(xlsx, sheet_name='Model Comparison'))   if 'Model Comparison'    in sheets else []
        tvt  = _clean(pd.read_excel(xlsx, sheet_name='Train vs Test'))      if 'Train vs Test'       in sheets else []
        bl   = _clean(pd.read_excel(xlsx, sheet_name='Baseline Results'))   if 'Baseline Results'    in sheets else []
        vraw = pd.read_excel(xlsx, sheet_name='Objective Validation')       if 'Objective Validation' in sheets else pd.DataFrame()
        val  = _clean(vraw)[0] if len(vraw) > 0 else {}

        COMPARISON = {'all_models': cmp, 'train_test': tvt, 'baseline': bl, 'validation': val}
    except Exception:
        pass


_load_comparison()


def _get_bo_params():
    if MODEL is None:
        return {}
    key_params = {
        'rf':  ['n_estimators', 'max_depth', 'min_samples_split', 'max_features'],
        'xgb': ['n_estimators', 'max_depth', 'learning_rate', 'subsample', 'colsample_bytree', 'reg_alpha'],
        'cat': ['learning_rate', 'depth', 'iterations', 'l2_leaf_reg'],
        'svr': ['C', 'epsilon', 'gamma'],
    }
    labels = {'rf': 'Random Forest', 'xgb': 'XGBoost', 'cat': 'CatBoost', 'svr': 'SVR'}
    result = {}
    try:
        for name, est in MODEL.named_estimators_.items():
            if name not in key_params:
                continue
            p = est.get_params()
            filtered = {k: (round(float(p[k]), 4) if isinstance(p[k], float) else p[k])
                        for k in key_params[name] if k in p}
            result[labels[name]] = filtered
    except Exception:
        pass
    return result


def pm25_to_aqi(pm25):
    pm25 = round(max(0.0, float(pm25)), 1)
    for c_lo, c_hi, a_lo, a_hi in AQI_BP:
        if c_lo <= pm25 <= c_hi:
            return round((a_hi - a_lo) / (c_hi - c_lo) * (pm25 - c_lo) + a_lo)
    return 500


def aqi_category(aqi):
    if aqi <=  50: return 'Good'
    if aqi <= 100: return 'Moderate'
    if aqi <= 150: return 'Unhealthy for Sensitive Groups'
    if aqi <= 200: return 'Unhealthy'
    if aqi <= 300: return 'Very Unhealthy'
    return 'Hazardous'


def _normalize_cols(df):
    """Map variant column names to standard names."""
    if all(isinstance(c, (int, np.integer)) for c in df.columns):
        n = len(df.columns)
        if n == 5:
            df.columns = REQUIRED_COLS
        elif n >= 6:
            df.columns = ['Date'] + REQUIRED_COLS + [f'_extra{i}' for i in range(n - 6)]
        return df

    aliases = {
        'Date':  {'date', 'datetime', 'timestamp', 'date/time'},
        'NO2':   {'no2', 'no2_hourly_avg1', 'no2_hourly'},
        'O3':    {'o3', 'o3_hourly_avg1', 'o3_hourly'},
        'CO':    {'co', 'co_hourly_avg1', 'co_hourly'},
        'PM10':  {'pm10', 'pm10_hourly_avg1', 'pm10_hourly'},
        'PM2.5': {'pm2.5', 'pm25', 'pm2.5_hourly_avg1', 'pm25_hourly'},
    }
    rename = {}
    for std, vs in aliases.items():
        for col in df.columns:
            stripped = col.strip()
            if stripped == std:
                break
            if stripped.lower() in vs:
                rename[col] = std
                break
    return df.rename(columns=rename)


def _read_file(file):
    """Try multiple read strategies until required columns are found."""
    data   = file.read()
    is_csv = file.filename.lower().endswith('.csv')

    def attempt(skip, header):
        buf = io.BytesIO(data)
        return (pd.read_csv if is_csv else pd.read_excel)(buf, skiprows=(skip or None), header=header)

    for skip, header in [(0, 0), (0, None), (2, 0), (1, 0)]:
        try:
            df = _normalize_cols(attempt(skip, header))
            if set(REQUIRED_COLS).issubset(df.columns):
                return df
        except Exception:
            pass

    raise ValueError('Cannot detect required columns (NO2, O3, CO, PM10, PM2.5). Please check your file format.')


def _build_features(df):
    """Clean and impute; generate lag-1 features only if the model requires them."""
    has_date = 'Date' in df.columns
    if has_date:
        df['Date'] = pd.to_datetime(df['Date'], errors='coerce')
        df = df.dropna(subset=['Date']).sort_values('Date').reset_index(drop=True)

    # Impute input features only; drop rows where PM2.5 (target) is missing
    for col in ['NO2', 'O3', 'CO', 'PM10']:
        df[col] = pd.to_numeric(df[col], errors='coerce').replace(0, np.nan)
        mean_val = df[col].mean()
        df[col] = df[col].fillna(mean_val if not np.isnan(mean_val) else 0.0)

    df['PM2.5'] = pd.to_numeric(df['PM2.5'], errors='coerce').replace(0, np.nan)
    df = df.dropna(subset=['PM2.5']).reset_index(drop=True)

    needs_lag = ALL_FEATURES and any('_lag1' in f for f in ALL_FEATURES)
    if needs_lag:
        for col in REQUIRED_COLS:
            df[f'{col}_lag1'] = df[col].shift(1)
        df = df.dropna(subset=[f'{col}_lag1' for col in REQUIRED_COLS]).reset_index(drop=True)

    return df, has_date


def _predict(raw_df):
    missing = [c for c in REQUIRED_COLS if c not in raw_df.columns]
    if missing:
        raise ValueError(f"Missing columns: {', '.join(missing)}. Required: {', '.join(REQUIRED_COLS)}")

    raw_cols = [c for c in (['Date'] + REQUIRED_COLS) if c in raw_df.columns]
    raw_prev = raw_df[raw_cols].head(10).copy()

    df, has_date = _build_features(raw_df.copy())

    if len(df) < 1:
        raise ValueError('No valid rows found after cleaning.')

    missing_features = [f for f in ALL_FEATURES if f not in df.columns]
    if missing_features:
        raise ValueError(f'Cannot build feature matrix — missing: {missing_features}')

    X = df[ALL_FEATURES].copy()
    scale_cols = [c for c in getattr(SCALER, 'feature_names_in_', ['PM10']) if c in X.columns]
    X[scale_cols] = SCALER.transform(X[scale_cols])

    n_prev    = min(10, len(X))
    prep_prev = X.head(n_prev).copy()
    prep_prev.insert(len(prep_prev.columns), 'PM2.5', df['PM2.5'].head(n_prev).values)

    preds  = MODEL.predict(X)
    actual = df['PM2.5'].values
    dates  = df['Date'].tolist() if has_date else None

    def _to_preview(d):
        cols = list(d.columns)
        rows = []
        for _, row in d.iterrows():
            r = []
            for v in row:
                if isinstance(v, pd.Timestamp):
                    r.append(str(v)[:19])
                elif isinstance(v, (float, np.floating)):
                    r.append(None if np.isnan(v) else round(float(v), 4))
                elif isinstance(v, (np.integer, np.int64)):
                    r.append(int(v))
                else:
                    try:
                        r.append(None if pd.isna(v) else v)
                    except Exception:
                        r.append(v)
            rows.append(r)
        return {'columns': cols, 'rows': rows}

    return (preds, actual, dates, df[REQUIRED_COLS].copy(),
            _to_preview(raw_prev), _to_preview(prep_prev))


@app.route('/model_comparison')
def model_comparison():
    _load_comparison()
    if COMPARISON is None:
        return jsonify({'error': 'No comparison data — run notebook Phase 11 first.'}), 404
    return jsonify(COMPARISON)


@app.route('/')
def index():
    return render_template('index.html', models_ok=(MODEL is not None))


@app.route('/predict', methods=['POST'])
def predict():
    if MODEL is None:
        return jsonify({'error': 'Models not loaded. Run the notebook (all phases) to generate saved_model/.'}), 500

    if 'file' not in request.files or not request.files['file'].filename:
        return jsonify({'error': 'No file provided.'}), 400

    file = request.files['file']

    try:
        raw = _read_file(file)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': f'Cannot read file: {e}'}), 400

    try:
        preds, actual, dates, pollutants, raw_preview, prep_preview = _predict(raw)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': f'Prediction failed: {e}'}), 500

    rmse = float(np.sqrt(mean_squared_error(actual, preds)))
    mae  = float(mean_absolute_error(actual, preds))
    r2   = float(r2_score(actual, preds))

    aqi_pred    = [pm25_to_aqi(v) for v in preds]
    aqi_actual  = [pm25_to_aqi(v) for v in actual]
    cats_pred   = [aqi_category(a) for a in aqi_pred]
    cats_actual = [aqi_category(a) for a in aqi_actual]
    cat_match   = float(np.mean(np.array(cats_pred) == np.array(cats_actual)) * 100)

    cat_counts = Counter(cats_pred)
    dominant   = max(cat_counts, key=cat_counts.get)

    if dates and isinstance(dates[0], pd.Timestamp):
        xlabels = [d.strftime('%Y-%m-%d %H:%M') for d in dates]
    else:
        xlabels = [f'#{i + 1}' for i in range(len(preds))]

    corr = pollutants.corr().round(3)

    ds_stats = {
        col: {
            'mean': round(float(pollutants[col].mean()), 4),
            'min':  round(float(pollutants[col].min()),  4),
            'max':  round(float(pollutants[col].max()),  4),
            'std':  round(float(pollutants[col].std()),  4),
        }
        for col in pollutants.columns
    }

    table = [
        {
            'idx':       i + 1,
            'date':      xlabels[i],
            'act_pm25':  round(float(actual[i]), 3),
            'pred_pm25': round(float(preds[i]),  3),
            'act_aqi':   aqi_actual[i],
            'pred_aqi':  aqi_pred[i],
            'category':  cats_pred[i],
            'color':     CAT_COLORS.get(cats_pred[i], '#7d8590'),
        }
        for i in range(min(25, len(preds)))
    ]

    return jsonify({
        'filename': file.filename,
        'metrics': {
            'rmse':      round(rmse, 4),
            'mae':       round(mae, 4),
            'r2':        round(r2, 4),
            'cat_match': round(cat_match, 1),
            'n_rows':    len(preds),
            'dominant':  dominant,
        },
        'charts': {
            'xlabels':    xlabels,
            'pm25_act':   [round(float(v), 3) for v in actual],
            'pm25_pred':  [round(float(v), 3) for v in preds],
            'aqi_act':    aqi_actual,
            'aqi_pred':   aqi_pred,
            'pie_labels': list(cat_counts.keys()),
            'pie_values': list(cat_counts.values()),
            'pie_colors': [CAT_COLORS.get(k, '#7d8590') for k in cat_counts],
            'corr_cols':  list(pollutants.columns),
            'corr_vals':  corr.values.tolist(),
        },
        'ds_stats': ds_stats,
        'table':    table,
        'pipeline': {
            'raw_preview':  raw_preview,
            'prep_preview': prep_preview,
            'features':     ALL_FEATURES or ['NO2', 'O3', 'CO', 'PM10'],
            'bo_params':    _get_bo_params(),
        },
    })


@app.route('/export', methods=['POST'])
def export():
    try:
        payload = request.get_json(force=True)
    except Exception:
        return jsonify({'error': 'Invalid request payload'}), 400

    filename  = payload.get('filename', 'prediction')
    metrics   = payload.get('metrics', {})
    charts    = payload.get('charts', {})
    ds_stats  = payload.get('ds_stats', {})

    xlabels   = charts.get('xlabels', [])
    pm25_act  = charts.get('pm25_act', [])
    pm25_pred = charts.get('pm25_pred', [])
    aqi_act   = charts.get('aqi_act', [])
    aqi_pred  = charts.get('aqi_pred', [])
    n = len(xlabels)

    if n == 0:
        return jsonify({'error': 'No prediction data to export'}), 400

    # ── Style helpers ──
    def _fill(hex6):
        return PatternFill('solid', fgColor='FF' + hex6.lstrip('#'))

    def _font(hex6, bold=False, sz=11):
        return Font(color='FF' + hex6.lstrip('#'), bold=bold, size=sz, name='Calibri')

    def _algn(h='center', wrap=False):
        return Alignment(horizontal=h, vertical='center', wrap_text=wrap)

    DARK_TEAL  = '1A6B57'
    MID_TEAL   = '2EA88A'
    LIGHT_TEAL = 'E8F5F1'
    WHITE      = 'FFFFFF'
    GREY_ALT   = 'F2F2F2'
    TEXT_DARK  = '222222'
    TEXT_MID   = '555555'
    TEXT_LIGHT = '999999'

    CAT_BG = {
        'Good':                           'D4EFDF',
        'Moderate':                       'FCF3CF',
        'Unhealthy for Sensitive Groups': 'FAE5D3',
        'Unhealthy':                      'FADBD8',
        'Very Unhealthy':                 'E8DAEF',
        'Hazardous':                      'F9EBEA',
    }
    CAT_FG = {
        'Good':                           '1E8449',
        'Moderate':                       '9A7D0A',
        'Unhealthy for Sensitive Groups': 'A04000',
        'Unhealthy':                      'CB4335',
        'Very Unhealthy':                 '76448A',
        'Hazardous':                      '922B21',
    }

    def _write(ws, row, col, value, bg=None, fg=None, bold=False, sz=11,
               h='left', wrap=False, merge_to=None):
        c = ws.cell(row=row, column=col, value=value)
        if bg:
            c.fill = _fill(bg)
        c.font = _font(fg or TEXT_DARK, bold=bold, sz=sz)
        c.alignment = _algn(h, wrap)
        if merge_to:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=row, end_column=merge_to)
        return c

    wb = openpyxl.Workbook()

    # ═══════════════════════════════════════════
    # SHEET 1: Summary
    # ═══════════════════════════════════════════
    ws1 = wb.active
    ws1.title = 'Summary'
    ws1.sheet_view.showGridLines = False

    _write(ws1, 1, 1, 'Air Quality PM2.5 — Prediction Report',
           bg=DARK_TEAL, fg=WHITE, bold=True, sz=14, h='center', merge_to=5)
    ws1.row_dimensions[1].height = 32

    _write(ws1, 2, 1,
           f"Source: {filename}   ·   Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
           bg=MID_TEAL, fg='E8F5F1', sz=10, h='center', merge_to=5)
    ws1.row_dimensions[2].height = 18
    ws1.row_dimensions[3].height = 8

    _write(ws1, 4, 1, '  PREDICTION METRICS',
           bg=DARK_TEAL, fg=WHITE, bold=True, sz=11, merge_to=5)
    ws1.row_dimensions[4].height = 22

    metrics_rows = [
        ('R² Score',              metrics.get('r2', '—'),
         'Goodness of fit · max 1.0 · higher = better'),
        ('RMSE',                  f"{metrics.get('rmse', '—')} μg/m³",
         'Root Mean Square Error'),
        ('MAE',                   f"{metrics.get('mae', '—')} μg/m³",
         'Mean Absolute Error'),
        ('AQI Category Match',    f"{metrics.get('cat_match', '—')}%",
         'Percentage of correct AQI category predictions'),
        ('Total Rows Predicted',  metrics.get('n_rows', '—'),
         'Number of data points processed'),
        ('Dominant AQI Category', metrics.get('dominant', '—'),
         'Most frequent predicted AQI category'),
    ]
    for ri, (label, value, note) in enumerate(metrics_rows, start=5):
        bg = LIGHT_TEAL if ri % 2 == 1 else WHITE
        _write(ws1, ri, 1, f'  {label}', bg=bg, fg=DARK_TEAL, bold=True)
        _write(ws1, ri, 2, value, bg=bg, fg=TEXT_DARK, bold=True, h='center')
        c = ws1.cell(row=ri, column=3, value=note)
        c.fill = _fill(bg); c.font = _font(TEXT_LIGHT); c.alignment = _algn('left')
        ws1.merge_cells(start_row=ri, start_column=3, end_row=ri, end_column=5)
        ws1.row_dimensions[ri].height = 20

    ws1.row_dimensions[11].height = 8
    _write(ws1, 12, 1, '  MODEL INFORMATION',
           bg=DARK_TEAL, fg=WHITE, bold=True, sz=11, merge_to=5)
    ws1.row_dimensions[12].height = 22

    model_rows = [
        ('Model Architecture', 'Stacking Ensemble'),
        ('Base Learners',      'Random Forest · XGBoost · CatBoost · SVR'),
        ('Meta Learner',       'XGBoost (200 est, depth=4, lr=0.05)'),
        ('Input Features',     f"{', '.join(ALL_FEATURES or ['NO2', 'O3', 'CO', 'PM10'])} ({len(ALL_FEATURES or [])} features)"),
        ('Target Variable',    'PM2.5 (μg/m³)'),
        ('AQI Standard',       'EPA PM2.5 Breakpoints'),
    ]
    for ri, (label, value) in enumerate(model_rows, start=13):
        bg = LIGHT_TEAL if ri % 2 == 1 else WHITE
        _write(ws1, ri, 1, f'  {label}', bg=bg, fg=DARK_TEAL, bold=True)
        c = ws1.cell(row=ri, column=2, value=value)
        c.fill = _fill(bg); c.font = _font(TEXT_DARK); c.alignment = _algn('left')
        ws1.merge_cells(start_row=ri, start_column=2, end_row=ri, end_column=5)
        ws1.row_dimensions[ri].height = 20

    ws1.row_dimensions[19].height = 8
    _write(ws1, 20, 1, '  AQI CATEGORY REFERENCE',
           bg=DARK_TEAL, fg=WHITE, bold=True, sz=11, merge_to=5)
    ws1.row_dimensions[20].height = 22

    for ci, hdr in enumerate(['Category', 'PM2.5 Range (μg/m³)', 'AQI Range', 'Description'], start=1):
        c = ws1.cell(row=21, column=ci, value=hdr)
        c.fill = _fill(MID_TEAL); c.font = _font(WHITE, bold=True, sz=10)
        c.alignment = _algn('center')
    ws1.merge_cells(start_row=21, start_column=4, end_row=21, end_column=5)
    ws1.row_dimensions[21].height = 20

    aqi_ref = [
        ('Good',                           '0.0 – 12.0',    '0 – 50',    'Air quality is satisfactory'),
        ('Moderate',                       '12.1 – 35.4',   '51 – 100',  'Acceptable; some may be sensitive'),
        ('Unhealthy for Sensitive Groups', '35.5 – 55.4',   '101 – 150', 'Sensitive groups may experience effects'),
        ('Unhealthy',                      '55.5 – 150.4',  '151 – 200', 'Everyone may experience health effects'),
        ('Very Unhealthy',                 '150.5 – 250.4', '201 – 300', 'Health alert: serious effects'),
        ('Hazardous',                      '250.5 – 500.4', '301 – 500', 'Emergency conditions'),
    ]
    for ri, (cat, pm_rng, aqi_rng, desc) in enumerate(aqi_ref, start=22):
        cbg = CAT_BG.get(cat, WHITE); cfg = CAT_FG.get(cat, TEXT_DARK)
        _write(ws1, ri, 1, cat, bg=cbg, fg=cfg, bold=True)
        _write(ws1, ri, 2, pm_rng, bg=cbg, fg=TEXT_DARK, h='center')
        _write(ws1, ri, 3, aqi_rng, bg=cbg, fg=TEXT_DARK, h='center')
        c = ws1.cell(row=ri, column=4, value=desc)
        c.fill = _fill(cbg); c.font = _font(TEXT_DARK); c.alignment = _algn('left')
        ws1.merge_cells(start_row=ri, start_column=4, end_row=ri, end_column=5)
        ws1.row_dimensions[ri].height = 20

    for col, w in zip('ABCDE', [30, 22, 16, 28, 8]):
        ws1.column_dimensions[col].width = w

    # ═══════════════════════════════════════════
    # SHEET 2: Prediction Results (all rows)
    # ═══════════════════════════════════════════
    ws2 = wb.create_sheet('Prediction Results')
    ws2.sheet_view.showGridLines = False

    p_cols   = ['#', 'Date / Index', 'Actual PM2.5\n(μg/m³)', 'Predicted PM2.5\n(μg/m³)',
                'Abs. Error\n(μg/m³)', 'Actual AQI', 'Predicted AQI', 'AQI Category']
    p_widths = [6, 22, 18, 20, 16, 13, 15, 32]

    ws2.merge_cells(f'A1:{get_column_letter(len(p_cols))}1')
    ws2['A1'] = f'Prediction Results — {filename}  ({n} rows)'
    ws2['A1'].fill = _fill(DARK_TEAL)
    ws2['A1'].font = _font(WHITE, bold=True, sz=13)
    ws2['A1'].alignment = _algn('center')
    ws2.row_dimensions[1].height = 28

    for ci, (hdr, w) in enumerate(zip(p_cols, p_widths), start=1):
        c = ws2.cell(row=2, column=ci, value=hdr)
        c.fill = _fill(MID_TEAL); c.font = _font(WHITE, bold=True, sz=10)
        c.alignment = _algn('center', wrap=True)
        ws2.column_dimensions[get_column_letter(ci)].width = w
    ws2.row_dimensions[2].height = 30

    for i in range(n):
        rn  = i + 3
        act = float(pm25_act[i])
        pre = float(pm25_pred[i])
        cat = aqi_category(int(aqi_pred[i]))
        err = round(abs(pre - act), 3)
        bg  = GREY_ALT if i % 2 == 1 else WHITE
        cbg = CAT_BG.get(cat, WHITE)
        cfg = CAT_FG.get(cat, TEXT_DARK)
        vals = [i + 1, xlabels[i], round(act, 3), round(pre, 3),
                err, int(aqi_act[i]), int(aqi_pred[i]), cat]
        for ci, val in enumerate(vals, start=1):
            c = ws2.cell(row=rn, column=ci, value=val)
            if ci == 8:
                c.fill = _fill(cbg); c.font = _font(cfg, bold=True, sz=10)
                c.alignment = _algn('left')
            else:
                c.fill = _fill(bg); c.font = _font(TEXT_DARK, sz=10)
                c.alignment = _algn('left' if ci == 2 else 'center')
        ws2.row_dimensions[rn].height = 17

    ws2.freeze_panes = 'A3'
    ws2.auto_filter.ref = f'A2:{get_column_letter(len(p_cols))}{n + 2}'

    # ═══════════════════════════════════════════
    # SHEET 3: Dataset Statistics
    # ═══════════════════════════════════════════
    ws3 = wb.create_sheet('Dataset Statistics')
    ws3.sheet_view.showGridLines = False

    ws3.merge_cells('A1:E1')
    ws3['A1'] = 'Dataset Statistics — Input Pollutants & Target Variable'
    ws3['A1'].fill = _fill(DARK_TEAL)
    ws3['A1'].font = _font(WHITE, bold=True, sz=13)
    ws3['A1'].alignment = _algn('center')
    ws3.row_dimensions[1].height = 28

    for ci, hdr in enumerate(['Pollutant', 'Mean', 'Minimum', 'Maximum', 'Std Dev'], start=1):
        c = ws3.cell(row=2, column=ci, value=hdr)
        c.fill = _fill(MID_TEAL); c.font = _font(WHITE, bold=True)
        c.alignment = _algn('center')
    ws3.row_dimensions[2].height = 22

    stat_units = {
        'NO2': 'NO2 (ppm)', 'O3': 'O3 (ppm)', 'CO': 'CO (mg/m³)',
        'PM10': 'PM10 (μg/m³)', 'PM2.5': 'PM2.5 (μg/m³)',
    }
    for ri, (col, stats) in enumerate(ds_stats.items(), start=3):
        bg    = LIGHT_TEAL if ri % 2 == 1 else WHITE
        label = stat_units.get(col, col)
        for ci, val in enumerate(
            [label, stats.get('mean'), stats.get('min'), stats.get('max'), stats.get('std')],
            start=1
        ):
            c = ws3.cell(row=ri, column=ci, value=val)
            c.fill = _fill(bg)
            c.font = _font(DARK_TEAL if ci == 1 else TEXT_DARK, bold=(ci == 1))
            c.alignment = _algn('left' if ci == 1 else 'center')
        ws3.row_dimensions[ri].height = 20

    for ci, w in enumerate([22, 14, 14, 14, 14], start=1):
        ws3.column_dimensions[get_column_letter(ci)].width = w
    ws3.freeze_panes = 'A3'

    # ── Stream back as .xlsx ──
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe = os.path.splitext(os.path.basename(filename))[0]
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'{safe}_report.xlsx',
    )


if __name__ == '__main__':
    app.run(debug=True, use_reloader=False)
